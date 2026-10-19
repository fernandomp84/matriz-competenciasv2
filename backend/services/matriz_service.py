"""
Servicio generador de Matriz_Competencias.csv a partir de JSON (v2).

La matriz es una lista plana de indicadores: una fila por ID Indicador con su
descripción (sin jerarquía matriz → metacompetencia → resultado de aprendizaje).
"""

import base64
import csv
import io
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


# ── Utilidades ─────────────────────────────────────────────────────────────────

# Espacios "raros" -> espacio normal; anchura cero -> se eliminan
_ESPACIOS_ESPECIALES = (" ", " ", " ", "\t")
_ANCHURA_CERO = ("​", "‌", "‍", "﻿")


def _normalizar_invisibles(texto: str) -> str:
    for ch in _ESPACIOS_ESPECIALES:
        texto = texto.replace(ch, " ")
    for ch in _ANCHURA_CERO:
        texto = texto.replace(ch, "")
    return texto


def limpiar(v) -> str:
    if not v:
        return ""
    t = str(v).replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
    t = _normalizar_invisibles(t)
    return re.sub(r" {2,}", " ", t).strip()


def _orden_id(codigo: str):
    """IDs numéricos en orden numérico; los demás al final, en orden alfabético."""
    return (0, int(codigo), "") if codigo.isdigit() else (1, 0, codigo)


# ── Generación de filas ─────────────────────────────────────────────────────────

def generar_filas(data: dict) -> tuple[list[dict], list[dict]]:
    """
    Una fila por ID Indicador con su descripción. Si un mismo ID llega con
    descripciones distintas (p. ej. un JSON editado a mano), se reporta y se omite.
    """
    descripciones: dict[str, set[str]] = {}
    cursos: dict[str, list[str]] = {}
    errores: list[dict] = []

    def err(id_siga: str, campo: str, mensaje: str):
        errores.append({"idCursoSiga": id_siga, "campo": campo, "error": mensaje})

    for id_siga, curso in data.items():
        for idx, comp in enumerate(curso.get("competencias", []), 1):
            try:
                ind = comp.get("indicadorDesempeno", {})
                codigo = limpiar(ind.get("codigo", ""))
                if not codigo:
                    err(id_siga, f"competencias[{idx}].indicadorDesempeno.codigo",
                        "Código de indicador vacío — fila ignorada")
                    continue
                desc = limpiar(ind.get("descripcion", ""))
                if not desc:
                    err(id_siga, f"competencias[{idx}].indicadorDesempeno.descripcion",
                        f"Indicador {codigo} sin descripción — fila ignorada")
                    continue
                descripciones.setdefault(codigo, set()).add(desc)
                cursos.setdefault(codigo, [])
                if id_siga not in cursos[codigo]:
                    cursos[codigo].append(id_siga)
            except Exception as exc:
                err(id_siga, f"competencias[{idx}]", f"Error inesperado: {exc}")

    filas: list[dict] = []
    for codigo in sorted(descripciones, key=_orden_id):
        descs = descripciones[codigo]
        if len(descs) > 1:
            err(", ".join(cursos[codigo][:5]), "indicadorDesempeno.descripcion",
                f"ID Indicador {codigo} tiene {len(descs)} descripciones distintas — indicador omitido")
            continue
        filas.append({"Short Code": codigo, "Description": next(iter(descs))})

    return filas, errores


# ── CSV ─────────────────────────────────────────────────────────────────────────

# Short Code = ID Indicador; Description = Indicador de desempeño
CAMPOS_CSV = ["Short Code", "Description"]


def generar_csv_string(filas: list[dict]) -> str:
    output = io.StringIO()
    writer = csv.DictWriter(
        output, fieldnames=CAMPOS_CSV, delimiter=";", quoting=csv.QUOTE_ALL
    )
    writer.writeheader()
    writer.writerows(filas)
    return output.getvalue()


# ── XLSX de errores ─────────────────────────────────────────────────────────────

def generar_errores_xlsx_b64(errores: list[dict]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Errores"

    thin = Side(style="thin", color="BFBFBF")
    brd = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", start_color="C00000")
    hdr_f = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    hdr_al = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = ["ID Curso SIGA", "Campo", "Error"]
    widths = [18, 40, 80]
    ws.row_dimensions[1].height = 28
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = hdr_f
        c.fill = hdr_fill
        c.alignment = hdr_al
        c.border = brd
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w

    data_font = Font(name="Arial", size=10)
    data_align = Alignment(vertical="top", wrap_text=True)
    alt_fill = PatternFill("solid", start_color="FFF2F2")

    for ri, e in enumerate(errores, 2):
        ws.row_dimensions[ri].height = 35
        fill = alt_fill if ri % 2 == 0 else PatternFill("solid", start_color="FFFFFF")
        for ci, key in enumerate(["idCursoSiga", "campo", "error"], 1):
            c = ws.cell(row=ri, column=ci, value=e.get(key, ""))
            c.font = data_font
            c.fill = fill
            c.alignment = data_align
            c.border = brd

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{len(errores) + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return base64.b64encode(buf.read()).decode("utf-8")


# ── Servicio principal ──────────────────────────────────────────────────────────

def procesar(json_data: dict) -> dict:
    data = json_data.get("data", json_data)
    filas, errores = generar_filas(data)
    csv_content = generar_csv_string(filas)
    errores_xlsx_b64 = generar_errores_xlsx_b64(errores)
    return {
        "filas": filas,
        "errores": errores,
        "csvContent": csv_content,
        "erroresXlsxB64": errores_xlsx_b64,
        "stats": {
            "totalFilas": len(filas),
            "totalErrores": len(errores),
        },
    }
