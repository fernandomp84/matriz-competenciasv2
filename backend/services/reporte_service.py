"""Reporte de errores (Errores_Megamatriz.xlsx): único punto de generación para la API y el script local."""
import io
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import megamatriz_service as mm

AZUL_ENCABEZADO = "00205B"
AZUL_SECCION = "DCE5F3"

TIPO_SIN_ID = "Curso omitido — tiene filas sin ID Indicador"
TIPO_CONFLICTO = "Curso omitido — usa indicadores en conflicto"
TIPO_DATOS = "Curso omitido — datos distintos entre programas"
TIPO_DEPTO = "Curso omitido — departamento sin equivalencia"


def tipo_error(e: dict) -> str:
    if "tiene filas de T_2 sin ID Indicador" in e["error"]:
        return TIPO_SIN_ID
    if e["error"].startswith("Fila sin ID Indicador"):
        return "Detalle — fila sin ID Indicador"
    if e["error"].startswith("Mismo indicador de desempeño con IDs distintos"):
        return "Advertencia — mismo texto con IDs distintos"
    if "usa indicadores" in e["error"]:
        return TIPO_CONFLICTO
    if "datos distintos entre programas" in e["error"]:
        return TIPO_DATOS
    if "Departamento" in e["error"]:
        return TIPO_DEPTO
    if e["error"].startswith("ID Indicador"):
        return "Indicador omitido — descripción/rúbrica distinta"
    return f"Otro ({e['hoja']})"


def _clave_orden(x: str):
    return (not x.isdigit(), int(x) if x.isdigit() else 0, x)


def _detalle_conflictos(t2: pd.DataFrame) -> list[dict]:
    cols = list(mm.CAMPOS_INDICADOR)
    d = t2.copy()
    d["ID Indicador"] = d["ID Indicador"].map(mm.id_indicador_a_str)
    for c in cols:
        d[c] = d[c].map(mm.limpiar_str)
    d["Fila Excel"] = d.index + 2
    conflictos = mm.analizar_indicadores(t2)
    det = []
    for id_ind in sorted(conflictos, key=_clave_orden):
        info = conflictos[id_ind]
        sub = d[d["ID Indicador"] == id_ind]
        for n, (_, grp) in enumerate(sub.groupby(cols, sort=False), 1):
            f0 = grp.iloc[0]
            det.append({
                "ID Indicador": id_ind, "Versión": n,
                "Campos que difieren": ", ".join(info["campos"]),
                "Filas Excel (hoja T_2)": ", ".join(map(str, grp["Fila Excel"])),
                "Programas (Source.Name)": ", ".join(sorted(grp["Source.Name"].astype(str).unique())),
                "Cursos": ", ".join(sorted(grp["Número de catálogo"].astype(str).unique())),
                **{c: f0[c] for c in cols},
            })
    return det


def _detalle_validaciones(t2: pd.DataFrame) -> tuple[list[dict], list[dict]]:
    sin_id = [{"Fila Excel (T_2)": f["fila"], "Número de catálogo": f["catalogo"],
               "Programa (Source.Name)": f["programa"], "Nivel de logro": f["nivel"],
               "Cód de Tributación": f["codTributacion"], "Indicador de desempeño": f["descripcion"]}
              for f in mm.filas_sin_id_indicador(t2)]
    mismo = []
    for g in mm.indicadores_mismo_texto(t2, set(mm.analizar_indicadores(t2))):
        for i in g["ids"]:
            d = g["detalle"][i]
            mismo.append({
                "Indicador de desempeño": g["descripcion"], "ID Indicador": i,
                "Filas T_2": ", ".join(map(str, d["filas"])), "Cursos": ", ".join(d["cursos"]),
                "Programas (Source.Name)": ", ".join(d["programas"]),
                "Resultado de aprendizaje": " | ".join(d["ras"]),
                "¿Rúbrica idéntica entre los IDs?": "Sí" if g["rubricaIdentica"] else "No",
                "¿Resultado de aprendizaje idéntico?": "Sí" if g["resultadoAprendizajeIdentico"] else "No",
            })
    return sin_id, mismo


def _resumen(resultado_mega: dict, resultado_matriz: dict, archivo: str, filas_err: list[dict]) -> list[dict]:
    st = resultado_mega["stats"]
    sm = resultado_matriz["stats"]
    # Un curso cuenta una vez por causa: los errores de T_1 se generan una vez por curso y causa
    cursos_por_tipo: dict[str, set] = {}
    for f in filas_err:
        cursos_por_tipo.setdefault(f["Tipo"], set()).add((f["Hoja"], f["Fila Excel"]))

    def n(tipo: str) -> int:
        return len(cursos_por_tipo.get(tipo, ()))

    fecha = datetime.now().strftime("%Y-%m-%d %H:%M")
    filas: list[dict] = []

    def titulo(t):
        if filas:
            filas.append({"Concepto": None, "Cantidad": None, "Ver hoja": None})
        filas.append({"Concepto": t, "Cantidad": None, "Ver hoja": None, "_titulo": True})

    def fila(c, q, v=None):
        # La referencia a una hoja solo tiene sentido si hay algo que ver (si es 0 la hoja no existe)
        filas.append({"Concepto": c, "Cantidad": q, "Ver hoja": v if q else None})

    titulo("ARCHIVO")
    fila(f"{archivo} — reporte generado {fecha}", None)
    titulo("RESULTADO")
    fila("Cursos en el Omnia (únicos)", st["cursosUnicosT1"])
    fila("Cursos procesados", st["procesados"])
    fila("Cursos omitidos", st["cursosOmitidos"])
    fila("Indicadores en la matriz", sm["totalFilas"])
    titulo("ERRORES (dejan cursos o indicadores por fuera)")
    fila("Cursos omitidos por filas sin ID Indicador", n(TIPO_SIN_ID),
         f'{st["filasSinId"]} filas — hoja "Filas sin ID"')
    fila("Cursos omitidos por indicadores en conflicto", n(TIPO_CONFLICTO), 'hoja "Indicadores en conflicto"')
    fila("Cursos omitidos por datos distintos entre programas", n(TIPO_DATOS), 'hoja "Errores"')
    fila("Cursos omitidos por departamento sin equivalencia", n(TIPO_DEPTO), 'hoja "Errores"')
    fila("Indicadores en conflicto", st["indicadoresOmitidos"], 'hoja "Indicadores en conflicto"')
    if sm["totalErrores"] > 0:
        fila("Errores al generar la matriz", sm["totalErrores"], 'hoja "Errores matriz"')
    titulo("ADVERTENCIAS (no omiten nada, solo revisar)")
    fila("Mismo texto con IDs distintos", st["advertencias"], 'hoja "Mismo texto - IDs distintos"')
    titulo("CONTROL")
    fila("Filas leídas en T_1 (una por programa)", st["totalFilasT1"])
    return filas


def _formatear(wb, hoja: str, anchos: list[int], datos: bool = True) -> None:
    ws = wb[hoja]
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=AZUL_ENCABEZADO)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for k, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(k)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if ws.cell(1, c.column).value == "ID Indicador":
                c.number_format = "@"  # Excel no debe quitar los ceros
    ws.freeze_panes = "A2" if hoja == "Resumen" else "B2"
    if datos:
        ws.auto_filter.ref = ws.dimensions


def generar_reporte_xlsx(sheets: dict, resultado_mega: dict, resultado_matriz: dict, archivo: str) -> bytes:
    hojas, _ = mm.resolver_hojas(sheets)
    t2 = mm.quitar_encabezados_repetidos(sheets[hojas["T2"]])

    filas_err = [{"Tipo": tipo_error(e), "Hoja": e["hoja"], "Fila Excel": e["fila"], "Detalle": e["error"]}
                 for e in resultado_mega["errores"] + resultado_mega["advertencias"]]
    resumen = _resumen(resultado_mega, resultado_matriz, archivo, filas_err)
    titulos = [i for i, f in enumerate(resumen) if f.pop("_titulo", False)]
    conflictos = _detalle_conflictos(t2)
    sin_id, mismo = _detalle_validaciones(t2)
    errores_matriz = resultado_matriz.get("errores") or []

    buf = io.BytesIO()
    anchos = {"Resumen": [58, 12, 42], "Errores": [48, 8, 11, 120]}
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        pd.DataFrame(resumen, columns=["Concepto", "Cantidad", "Ver hoja"]).to_excel(
            w, sheet_name="Resumen", index=False)
        (pd.DataFrame(filas_err, columns=["Tipo", "Hoja", "Fila Excel", "Detalle"])
         .sort_values(["Hoja", "Fila Excel"]).to_excel(w, sheet_name="Errores", index=False))
        if conflictos:
            pd.DataFrame(conflictos).to_excel(w, sheet_name="Indicadores en conflicto", index=False)
            anchos["Indicadores en conflicto"] = [14, 9, 40, 22, 30, 25, 60, 50, 50, 50, 50]
        if sin_id:
            pd.DataFrame(sin_id).to_excel(w, sheet_name="Filas sin ID", index=False)
            anchos["Filas sin ID"] = [14, 16, 22, 14, 36, 90]
        if mismo:
            pd.DataFrame(mismo).to_excel(w, sheet_name="Mismo texto - IDs distintos", index=False)
            anchos["Mismo texto - IDs distintos"] = [60, 14, 24, 16, 30, 60, 16, 16]
        if errores_matriz:
            pd.DataFrame(errores_matriz).to_excel(w, sheet_name="Errores matriz", index=False)
            anchos["Errores matriz"] = [30, 14, 100, 30, 30]

    buf.seek(0)
    wb = load_workbook(buf)
    for nombre, a in anchos.items():
        _formatear(wb, nombre, a, datos=nombre != "Resumen")
    ws = wb["Resumen"]
    for i in titulos:  # fila de datos i -> fila Excel i + 2
        for c in ws[i + 2]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", start_color=AZUL_SECCION)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
