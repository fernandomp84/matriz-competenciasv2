"""
Genera en local los 4 archivos de salida a partir de un Omnia, con la misma lógica del backend v2:

    Salida/salida.json
    Salida/Matriz_Competencias.csv
    Salida/Inconsistencias_ID_Indicador.xlsx
    Salida/Errores_Megamatriz.xlsx

Uso (desde la carpeta Matriz 2):
    python scripts/generar_salidas.py "Entrada/Omnia2.xlsx"

La Tabla de equivalencias se toma de Entrada/Tabla de equivalencias.xlsx.
"""
import json
import os
import sys
from collections import Counter
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from backend.services import matriz_service  # noqa: E402
from backend.services import megamatriz_service as mm  # noqa: E402

SALIDA = os.path.join(ROOT, "Salida")
EQUIV = os.path.join(ROOT, "Entrada", "Tabla de equivalencias.xlsx")


def formatear(path: str, anchos: dict[str, list[int]]) -> None:
    """Encabezados en azul institucional, ajuste de texto, filtros e ID Indicador como texto."""
    wb = load_workbook(path)
    for nombre, widths in anchos.items():
        ws = wb[nombre]
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", start_color="00205B")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for k, w in enumerate(widths):
            ws.column_dimensions[chr(65 + k)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if ws.cell(1, c.column).value == "ID Indicador":
                    c.number_format = "@"  # Excel no debe quitar los ceros
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def reporte_inconsistencias(main: bytes, path: str) -> int:
    sheets = mm.leer_hojas(main)
    hojas, _ = mm.resolver_hojas(sheets)
    t2 = mm.quitar_encabezados_repetidos(sheets[hojas["T2"]])
    cols = list(mm.CAMPOS_INDICADOR)
    d = t2.copy()
    d["ID Indicador"] = d["ID Indicador"].map(mm.id_indicador_a_str)
    for c in cols:
        d[c] = d[c].map(mm.limpiar_str)
    d["Fila Excel"] = d.index + 2

    conflictos = mm.analizar_indicadores(t2)
    res, det = [], []
    for id_ind in sorted(conflictos, key=lambda x: (not x.isdigit(), int(x) if x.isdigit() else 0, x)):
        info = conflictos[id_ind]
        sub = d[d["ID Indicador"] == id_ind]
        variantes = sub.groupby(cols, sort=False)
        res.append({
            "ID Indicador": id_ind,
            "Campos que difieren": ", ".join(info["campos"]),
            "Versiones distintas": variantes.ngroups,
            "Filas en T_2": len(sub),
            "Cursos": ", ".join(sorted(sub["Número de catálogo"].astype(str).unique())),
        })
        for n, (_, grp) in enumerate(variantes, 1):
            f0 = grp.iloc[0]
            det.append({
                "ID Indicador": id_ind, "Versión": n,
                "Campos que difieren": ", ".join(info["campos"]),
                "Filas Excel (hoja T_2)": ", ".join(map(str, grp["Fila Excel"])),
                "Programas (Source.Name)": ", ".join(sorted(grp["Source.Name"].astype(str).unique())),
                "Cursos": ", ".join(sorted(grp["Número de catálogo"].astype(str).unique())),
                **{c: f0[c] for c in cols},
            })
    with pd.ExcelWriter(path) as w:
        pd.DataFrame(res, columns=["ID Indicador", "Campos que difieren", "Versiones distintas",
                                   "Filas en T_2", "Cursos"]).to_excel(w, sheet_name="Resumen", index=False)
        pd.DataFrame(det).to_excel(w, sheet_name="Detalle", index=False)
    formatear(path, {"Resumen": [14, 45, 12, 12, 60], "Detalle": [14, 9, 40, 22, 30, 25, 60, 50, 50, 50, 50]})
    return len(res)


def tipo_error(e: dict) -> str:
    if "tiene filas de T_2 sin ID Indicador" in e["error"]:
        return "Curso omitido — tiene filas sin ID Indicador"
    if e["error"].startswith("Fila sin ID Indicador"):
        return "Detalle — fila sin ID Indicador"
    if e["error"].startswith("Mismo indicador de desempeño con IDs distintos"):
        return "Advertencia — mismo texto con IDs distintos"
    if "usa indicadores" in e["error"]:
        return "Curso omitido — usa indicadores en conflicto"
    if "datos distintos entre programas" in e["error"]:
        return "Curso omitido — datos distintos entre programas"
    if "Departamento" in e["error"]:
        return "Curso omitido — departamento sin equivalencia"
    if e["error"].startswith("ID Indicador"):
        return "Indicador omitido — descripción/rúbrica distinta"
    return f"Otro ({e['hoja']})"


def detalle_validaciones(main: bytes) -> tuple[list[dict], list[dict]]:
    """Detalle para las hojas 'Filas sin ID' y 'Mismo texto - IDs distintos'."""
    sheets = mm.leer_hojas(main)
    hojas, _ = mm.resolver_hojas(sheets)
    t2 = mm.quitar_encabezados_repetidos(sheets[hojas["T2"]])
    sin_id = [{"Fila Excel (T_2)": f["fila"], "Número de catálogo": f["catalogo"],
               "Programa (Source.Name)": f["programa"], "Nivel de logro": f["nivel"],
               "Cód de Tributación": f["codTributacion"], "Indicador de desempeño": f["descripcion"]}
              for f in mm.filas_sin_id_indicador(t2)]
    mismo = []
    for g in mm.indicadores_mismo_texto(t2, set(mm.analizar_indicadores(t2))):
        for i in g["ids"]:
            d = g["detalle"][i]
            mismo.append({
                "Indicador de desempeño": g["descripcion"], "ID Indicador": i,
                "Filas T_2": ", ".join(map(str, d["filas"])), "Cursos": ", ".join(d["cursos"]),
                "Programas (Source.Name)": ", ".join(d["programas"]),
                "Resultado de aprendizaje": " | ".join(d["ras"]),
                "¿Rúbrica idéntica entre los IDs?": "Sí" if g["rubricaIdentica"] else "No",
                "¿Resultado de aprendizaje idéntico?": "Sí" if g["resultadoAprendizajeIdentico"] else "No",
            })
    return sin_id, mismo


def reporte_errores(r: dict, m: dict, archivo: str, path: str, detalle: tuple) -> Counter:
    sin_id, mismo = detalle
    filas = [{"Tipo": tipo_error(e), "Hoja": e["hoja"], "Fila Excel": e["fila"], "Detalle": e["error"]}
             for e in r["errores"] + r["advertencias"]]
    tipos = Counter(f["Tipo"] for f in filas)
    st = r["stats"]
    resumen = [{"Tipo": f"Archivo: {archivo}", "Cantidad": None}]
    resumen += [{"Tipo": t, "Cantidad": n} for t, n in tipos.items()]
    resumen += [
        {"Tipo": "Filas en T_1 (una por programa)", "Cantidad": st["totalFilasT1"]},
        {"Tipo": "Cursos únicos en T_1 (ID SIGA)", "Cantidad": st["cursosUnicosT1"]},
        {"Tipo": "Cursos únicos procesados (en el JSON)", "Cantidad": st["procesados"]},
        {"Tipo": "Indicadores en la matriz", "Cantidad": m["stats"]["totalFilas"]},
        {"Tipo": "Errores al generar la matriz", "Cantidad": m["stats"]["totalErrores"]},
    ]
    with pd.ExcelWriter(path) as w:
        pd.DataFrame(resumen).to_excel(w, sheet_name="Resumen", index=False)
        (pd.DataFrame(filas, columns=["Tipo", "Hoja", "Fila Excel", "Detalle"])
         .sort_values(["Hoja", "Fila Excel"]).to_excel(w, sheet_name="Errores", index=False))
        if m["errores"]:
            pd.DataFrame(m["errores"]).to_excel(w, sheet_name="Errores matriz", index=False)
        if sin_id:
            pd.DataFrame(sin_id).to_excel(w, sheet_name="Filas sin ID", index=False)
        if mismo:
            pd.DataFrame(mismo).to_excel(w, sheet_name="Mismo texto - IDs distintos", index=False)
    anchos = {"Resumen": [55, 12], "Errores": [48, 8, 11, 120]}
    if sin_id:
        anchos["Filas sin ID"] = [14, 16, 22, 14, 36, 90]
    if mismo:
        anchos["Mismo texto - IDs distintos"] = [60, 14, 24, 16, 30, 60, 16, 16]
    formatear(path, anchos)
    return tipos


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    omnia = os.path.abspath(sys.argv[1])
    os.makedirs(SALIDA, exist_ok=True)
    datos = open(omnia, "rb").read()
    equiv = open(EQUIV, "rb").read()
    fecha = datetime.fromtimestamp(os.path.getmtime(omnia)).strftime("%Y-%m-%d %H:%M")
    archivo = f"{os.path.basename(omnia)} ({fecha})"

    r = mm.procesar(datos, equiv, os.path.basename(omnia))
    m = matriz_service.procesar(r["resultado"])
    with open(os.path.join(SALIDA, "salida.json"), "w", encoding="utf-8") as f:
        json.dump(r["resultado"], f, ensure_ascii=False, indent=2)
    with open(os.path.join(SALIDA, "Matriz_Competencias.csv"), "w", encoding="utf-8-sig", newline="") as f:
        f.write(m["csvContent"])
    n_inc = reporte_inconsistencias(datos, os.path.join(SALIDA, "Inconsistencias_ID_Indicador.xlsx"))
    tipos = reporte_errores(r, m, archivo, os.path.join(SALIDA, "Errores_Megamatriz.xlsx"),
                            detalle_validaciones(datos))

    st = r["stats"]
    print(f"Archivo: {archivo}")
    print(f"Cursos únicos: {st['cursosUnicosT1']} | procesados: {st['procesados']} | omitidos: {st['cursosOmitidos']}")
    print(f"Indicadores en conflicto: {n_inc} | indicadores en la matriz: {m['stats']['totalFilas']}")
    print("Errores y advertencias:", dict(tipos))
    print("Archivos generados en", SALIDA)


if __name__ == "__main__":
    main()
